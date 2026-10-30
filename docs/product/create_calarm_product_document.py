from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\sondh\Documents\SmartWorkAlarm")
OUTPUT = ROOT / "artifacts" / "Calarm_제품_기능_마케팅_유료화_정의서.docx"
ICON = ROOT / "app" / "src" / "main" / "res" / "drawable-nodpi" / "ic_launcher_calarm.jpg"

NAVY = "102846"
DEEP_NAVY = "07111F"
BLUE = "315DA8"
MINT = "36BFA7"
PALE_BLUE = "EAF2FC"
PALE_MINT = "E9F8F4"
PALE_GRAY = "F3F5F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(table):
    repeat_header(table.rows[0])


def set_run_font(run, name="Malgun Gothic", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, name="Malgun Gothic", size=None, bold=None, color=None):
    for run in paragraph.runs:
        set_run_font(run, name, size, bold, color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r_fonts.set(qn("w:ascii"), "Malgun Gothic")
    r_fonts.set(qn("w:hAnsi"), "Malgun Gothic")
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    set_run_font(run, size=9, color=MID_GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Malgun Gothic"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, color in (("Heading 1", 19, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11, DEEP_NAVY)):
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(13 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.add_run("Calarm 제품 정의서    ")
    set_paragraph_font(footer, size=9, color=MID_GRAY)
    add_page_number(footer)


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(item))


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(f"{index}.  {item}"))


def callout(doc, title, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 170, 190, 170, 190)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    set_run_font(p2.add_run(text), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def make_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(header), size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Cm(widths[idx])
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_row(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(str(value)), size=font_size)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def page_break(doc):
    doc.add_page_break()


doc = Document()
configure_document(doc)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(16)
p.add_run().add_picture(str(ICON), width=Inches(1.45))

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p.add_run("Calarm 제품 기능 마케팅 유료화 정의서")
p_ppr = p._element.get_or_add_pPr()
p_border = p_ppr.find(qn("w:pBdr"))
if p_border is not None:
    p_ppr.remove(p_border)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("불규칙한 출근 일정에 맞춰 모닝콜을 자동 선택하는 캘린더 기반 알람"), size=13, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(28)
set_run_font(p.add_run("제품 기준 버전 1.10.3\n작성일 2026년 9월 5일\n문서 성격 제품 및 사업 통합 정의서"), size=10, color=MID_GRAY)

doc.add_paragraph().paragraph_format.space_before = Pt(34)
callout(
    doc,
    "한 문장 정의",
    "근무 일정은 캘린더에 기록하고 기상 시간은 한 번만 연결하면 Calarm이 매일 필요한 모닝콜을 자동으로 선택한다.",
    fill=PALE_MINT,
    accent=MINT,
)

page_break(doc)

# Executive summary
heading(doc, "문서 목적", 1)
body(doc, "이 문서는 Calarm의 제품 목적과 현재 기능을 정의하고 개인 사용자와 회사 사용자의 가치를 설명하며 무료 및 유료 기능의 경계와 출시 이후 마케팅 및 수익화 방안을 제시한다. 현재 구현된 기능과 향후 사업 계획을 구분해 제품 개발과 소개 자료와 의사결정의 공통 기준으로 사용한다.")

heading(doc, "핵심 결론", 2)
make_table(doc, ["구분", "정의"], [
    ("제품 본질", "캘린더 앱이 아니라 불규칙한 출근 시간을 자동으로 따라가는 모닝콜 앱"),
    ("무료 가치", "반복 알람과 캘린더 일정에 따른 자동 건너뛰기로 핵심 편리함을 충분히 체험"),
    ("개인 유료", "교대 그룹 생성부터 유료로 제공하고 최초에는 평생 이용권으로 단순하게 출시"),
    ("회사 가치", "사업주가 공용 캘린더에 근무표를 입력하고 직원은 자신의 기상 시간만 설정"),
    ("회사 수익화", "초기에는 공유 캘린더 기반으로 검증하고 중앙 관리 기능이 생길 때 월 구독으로 확장"),
    ("최우선 품질", "알람 정확성 잠금화면 접근 권한 진단 캘린더 조회 실패 시 안전 동작"),
], widths=[3.2, 12.8], font_size=9.2)

callout(doc, "제품 약속", "사용자는 매일 알람을 다시 맞추지 않아도 된다. 캘린더의 근무 일정만 정확하면 Calarm이 그날 필요한 알람을 선택한다.")

page_break(doc)
heading(doc, "제품 상태 구분", 2)
bullets(doc, [
    "현재 제공 기능은 Calarm 1.10.3 코드와 앱 도움말을 기준으로 작성했다.",
    "현재 앱에는 결제나 기능 잠금이 적용되어 있지 않다. 무료와 유료 구분은 향후 도입할 상품 정책이다.",
    "회사 공용 캘린더 활용은 현재 기능으로 가능하지만 직원 관리 대시보드와 중앙 배포 기능은 향후 범위다.",
])

# Product purpose
heading(doc, "제품 목적", 1)
heading(doc, "출발점", 2)
body(doc, "Calarm의 초기 목적은 알람이며 그중에서도 출근을 위한 모닝콜의 비중이 가장 크다. 일반 알람 앱은 정해진 요일과 시간의 반복에는 편리하지만 교대근무와 변동 근무처럼 출근 시간이 자주 달라지는 생활에서는 사용자가 계속 알람을 수정하거나 켜고 꺼야 한다.")
body(doc, "캘린더 앱 자체 알림은 일정을 알려주는 데 적합하지만 사용자가 일정마다 기상 시간을 역산해 알림을 별도로 입력해야 하고 강한 알람음과 잠금화면 해제와 다시 울림 같은 모닝콜 경험이 부족하다. Calarm은 캘린더를 일정 입력 도구로 활용하면서 실제 기상 행동은 전문 알람 화면과 소리로 처리한다.")

heading(doc, "해결하려는 문제", 2)
make_table(doc, ["기존 방식", "사용자 불편", "Calarm의 해결"], [
    ("일반 반복 알람", "근무가 바뀔 때마다 여러 알람을 직접 켜고 꺼야 함", "캘린더의 근무 형태를 읽어 필요한 알람만 자동 선택"),
    ("캘린더 일정 알림", "매 일정마다 기상 시간을 계산하고 알림을 입력해야 함", "근무 형태별 기상 시간을 한 번 저장하고 일정 제목만 등록"),
    ("휴가 전날 수동 해제", "알람을 끄는 것을 잊거나 휴가 뒤 다시 켜는 것을 잊음", "휴가 키워드 날짜만 건너뛰고 다음 반복일부터 자동 복귀"),
    ("회사 근무표 전달", "직원이 근무표를 보고 각자 알람을 다시 설정", "공용 캘린더 스케줄을 직원 Calarm이 읽고 개인 기상 시간 적용"),
], widths=[3.2, 6.4, 6.4], font_size=8.8)

heading(doc, "제품 원칙", 2)
numbered(doc, [
    "알람 신뢰성이 기능 수보다 우선한다.",
    "사용자는 캘린더에 근무 형태만 기록하고 기상 시간은 반복 입력하지 않는다.",
    "자동 판단 결과를 22시 알림과 다음 알람 화면으로 미리 확인할 수 있어야 한다.",
    "캘린더를 읽지 못할 때는 조용히 실패하지 않고 안전 규칙에 따라 동작한다.",
    "무료 기능만으로도 제품의 핵심 가치를 경험하게 하고 복잡한 교대 자동화에서 유료 가치를 만든다.",
    "초기에는 서버 없이 휴대폰과 캘린더 데이터를 중심으로 운영 비용과 개인정보 부담을 낮춘다.",
])

# Personas
heading(doc, "핵심 사용자", 1)
heading(doc, "개인 사용자", 2)
make_table(doc, ["사용자", "상황", "핵심 요구", "제품 가치"], [
    ("교대근무자", "주간 야간 새벽 근무가 반복 또는 수시 변경", "오늘 근무에 맞는 알람 하나만 정확히 울림", "근무 형태별 기상 시간을 한 번만 설정"),
    ("불규칙 출근 근로자", "출근 시간과 요일이 일정하지 않음", "캘린더 일정과 알람을 자동 연결", "매일 알람 재설정 제거"),
    ("휴가와 출장이 잦은 직장인", "쉬는 날에도 평일 알람이 울릴 수 있음", "특정 일정 날짜만 자동 건너뛰기", "끄고 다시 켜는 실수 방지"),
], widths=[3.0, 4.4, 4.6, 4.4], font_size=8.4)

heading(doc, "회사 사용자", 2)
make_table(doc, ["사용자", "역할", "핵심 요구", "제품 가치"], [
    ("사업주", "불규칙 근무자를 고용하고 근무표 운영", "직원이 회사 스케줄을 따라 알람을 설정하도록 단순화", "출근 준비 실수와 반복 안내 감소"),
    ("스케줄 담당자", "팀별 주간 야간 근무를 배정", "익숙한 Google Calendar에 근무 형태 입력", "별도 관리 시스템 없이 빠른 도입"),
    ("직원", "회사 근무표를 보고 자신의 기상 시간을 결정", "회사 일정은 자동 반영하고 기상 시간은 개인이 통제", "개인 생활과 회사 일정의 역할 분리"),
], widths=[3.0, 4.4, 4.6, 4.4], font_size=8.4)

callout(doc, "회사 버전의 핵심", "회사는 언제 출근해야 하는지를 제공하고 직원은 몇 시에 일어날지를 결정한다. Calarm은 두 정보를 결합해 실제 모닝콜을 만든다.", fill=PALE_MINT, accent=MINT)

heading(doc, "우선 공략 시장", 2)
bullets(doc, [
    "제조 생산 물류 보안 의료 숙박 운송 시설관리처럼 교대 근무가 흔한 직군",
    "매장별 오픈 마감 시간이 바뀌는 유통 및 서비스업",
    "현장 배치와 호출 시간이 자주 달라지는 단기 및 계약 근로자",
    "정형화된 인사 시스템을 도입하기 어려운 소규모 사업장",
])

# Value proposition
heading(doc, "가치 제안과 차별점", 1)
heading(doc, "핵심 가치 문장", 2)
callout(doc, "사용자용", "근무표만 캘린더에 적어 두세요. 오늘 필요한 모닝콜은 Calarm이 고릅니다.")
callout(doc, "사업주용", "공용 캘린더에 근무 형태를 등록하면 직원은 자신의 기상 시간만 정하면 됩니다.", fill=PALE_MINT, accent=MINT)

heading(doc, "일반 알람과 캘린더 알림 대비", 2)
make_table(doc, ["평가 기준", "일반 알람", "캘린더 알림", "Calarm"], [
    ("불규칙 근무 대응", "수동 켜기 끄기", "일정마다 별도 알림", "근무 형태 자동 선택"),
    ("모닝콜 경험", "강함", "제한적", "전체화면 소리 진동 다시 울림"),
    ("휴가 자동 처리", "대부분 수동", "일정 알림만 중단", "해당 날짜 알람 자동 건너뛰기"),
    ("회사 근무표 활용", "근무표를 보고 재설정", "일정 공유 가능", "공유 일정과 개인 기상 시간 결합"),
    ("사전 확인", "다음 알람 표시", "일정 표시", "22시 상태 변경 알림과 다음 알람"),
], widths=[3.3, 4.0, 4.0, 5.1], font_size=8.4)

page_break(doc)

# Feature definitions
heading(doc, "현재 기능 정의", 1)
body(doc, "다음 표는 Calarm 1.10.3에서 제공하는 기능을 사용자 관점으로 정의한다. 결제 기능은 아직 없으므로 교대 그룹도 현재 빌드에서는 제한 없이 사용할 수 있다.")

make_table(doc, ["기능", "정의", "주요 규칙", "상태"], [
    ("반복 알람", "알람 이름 시간 반복 요일을 지정", "선택 요일마다 반복하며 현재 알람 해제 후에도 다음 반복 유지", "구현"),
    ("다중 알람", "사용자가 여러 개의 일반 알람 생성", "각 알람은 독립적인 시간 요일 소리와 캘린더 조건 보유", "구현"),
    ("캘린더 자동 건너뛰기", "일정 제목에 지정 키워드가 포함된 날짜의 알람을 건너뜀", "쉼표로 여러 키워드 입력 가능하고 선택한 Google 캘린더만 조회", "구현"),
    ("휴가 자동 처리", "휴가 일정 기간의 모닝 알람을 중지", "하루 일정은 하루만 중지하고 종료 후 다음 반복일부터 자동 복귀", "구현"),
    ("공휴일 자동 끄기", "법정공휴일에 해당 알람을 건너뜀", "대한민국 2026년부터 2031년 내장 데이터와 로컬 캘린더 보완", "구현"),
    ("교대 그룹", "주간 야간 새벽 등 여러 근무 알람을 하나의 그룹으로 관리", "일정 제목 전체와 별칭이 정확히 일치하는 근무 형태 하나만 울림", "구현 현재 무료"),
    ("교대 충돌 방지", "같은 날 여러 근무 형태가 일치하면 그룹 알람 중지", "22시 확인에서 사용자에게 근무표 확인 경고", "구현"),
    ("비상 기본 알람", "캘린더 조회 실패 시 사용할 교대 알람 지정", "조회 실패 때만 사용하며 미지정 시 교대 그룹 건너뜀", "구현"),
    ("22시 다음 날 확인", "매일 22시에 다음 날 상태를 확인", "켜짐 꺼짐 또는 근무 형태 변경이 있을 때만 시스템 알림", "구현"),
    ("알람 30분 전 확인", "실제로 울릴 알람에 한해 오늘 끌지 질문", "상단 알림에서 끄기 또는 유지 선택 가능", "구현"),
    ("최종 재확인", "22시 이후 일정 변경을 반영", "실제 알람 시각에 캘린더를 다시 읽어 울림 여부 최종 판단", "구현"),
    ("알람음", "시스템 알람음 및 앱 내 추천 소리 선택", "추천 소리는 서버 없이 APK에 포함", "구현"),
    ("점점 크게", "알람 음량을 점진적으로 증가", "알람별 선택", "구현"),
    ("다시 울림", "5분부터 30분 사이 재알림", "알람별 사용 여부 선택하고 알람 화면을 3초간 길게 눌러 실행", "구현"),
    ("잠금화면 알람", "잠긴 화면 위에 전용 알람 화면 표시", "전체화면 권한과 제조사 설정의 영향을 진단", "구현"),
    ("밀어서 해제", "중앙 원을 어느 방향이든 밀어 알람 해제", "성공 시 민 방향으로 화면 밖 이동하고 미달 시 중앙 복귀", "구현"),
    ("다음 알람 표시", "홈 화면에서 다음 실행 예정 알람 제공", "아래로 당겨 새로고침 가능하고 저장 후 즉시 재계산", "구현"),
    ("캘린더 선택", "Google 계정별 캘린더를 선택", "하나 이상 선택해야 연결 상태이며 Outlook은 아직 미지원", "구현"),
    ("권한 및 진단", "알림 정확한 알람 전체화면 배터리 설정 확인", "필수 권한이 없으면 앱 실행과 편집 시 지속 안내", "구현"),
], widths=[3.0, 5.3, 6.3, 1.8], font_size=7.7)

heading(doc, "캘린더 판단 규칙", 2)
bullets(doc, [
    "일반 알람의 키워드는 일정 제목에 포함되면 일치한다.",
    "교대 그룹의 근무 형태는 앞뒤 공백과 대소문자를 제외하고 일정 제목 전체가 등록 별칭과 같아야 한다.",
    "휴가와 공휴일 조건은 교대 근무 선택보다 우선한다.",
    "일치하는 근무 형태가 없거나 서로 다른 근무 형태가 동시에 일치하면 교대 그룹은 울리지 않는다.",
    "일반 알람은 캘린더 조회 실패 시 놓치지 않도록 울리고 교대 그룹은 비상 기본 알람 설정을 따른다.",
])

# Journeys
heading(doc, "핵심 사용 흐름", 1)
heading(doc, "개인 무료 사용 흐름", 2)
numbered(doc, [
    "사용자가 평일 반복 모닝 알람을 만들고 시간과 요일을 저장한다.",
    "캘린더 자동 건너뛰기를 켜고 휴가 또는 사용자가 원하는 키워드를 입력한다.",
    "Google Calendar에 해당 키워드가 포함된 일정을 등록한다.",
    "Calarm이 그 날짜의 알람만 건너뛰고 다음 반복일에는 자동으로 다시 울린다.",
])

heading(doc, "개인 유료 교대근무 흐름", 2)
numbered(doc, [
    "사용자가 교대 그룹을 만들고 주간 야간 새벽 근무 형태를 추가한다.",
    "각 근무 형태에 알람 시간과 정확한 캘린더 제목 또는 별칭을 연결한다.",
    "사용자는 앞으로 캘린더에 근무 형태 제목만 날짜별로 입력한다.",
    "Calarm이 해당 날짜에 일치하는 근무 알람 하나만 선택하고 변경 시 22시에 알린다.",
])

heading(doc, "회사 공용 캘린더 흐름", 2)
numbered(doc, [
    "회사는 공유용 Google Calendar를 만들고 스케줄 담당자에게 편집 권한을 준다.",
    "직원에게는 일정 세부정보를 볼 수 있는 최소 권한으로 캘린더를 공유한다.",
    "담당자는 날짜별 근무 형태를 일정 제목으로 등록한다.",
    "직원은 휴대폰에서 회사 캘린더를 동기화하고 Calarm에서 해당 캘린더를 선택한다.",
    "직원은 근무 형태별로 자신의 기상 시간만 지정한다.",
    "회사 근무표가 바뀌면 직원 Calarm이 다음 확인 시 변경된 알람을 자동 선택한다.",
])

callout(doc, "보안 원칙", "회사 계정의 아이디와 비밀번호를 전 직원이 공유하지 않는다. 회사가 소유한 별도 캘린더를 직원 개인 계정 또는 조직 그룹에 읽기 권한으로 공유하는 방식을 권장한다.", fill="FFF4E5", accent="B54708")

heading(doc, "현재 회사 활용의 제한", 2)
bullets(doc, [
    "직원별 일정이 한 캘린더에 섞이면 직원별 전체 제목 또는 팀별 캘린더 분리가 필요하다.",
    "다른 직원의 휴가라는 제목까지 함께 읽으면 내 알람에도 영향을 줄 수 있으므로 직원 또는 팀별 휴가 캘린더를 분리하는 편이 안전하다.",
    "현재는 회사 관리자가 직원의 설치 상태 알람 설정 상태 또는 실제 알람 성공을 중앙에서 확인할 수 없다.",
    "Google Calendar 동기화 지연과 Android 절전 정책은 회사가 통제할 수 없으므로 초기 도입 시 기기별 점검 절차가 필요하다.",
])

page_break(doc)

# Monetization
heading(doc, "무료 및 유료 상품 정의", 1)
body(doc, "유료화의 기준은 캘린더 연동 자체가 아니라 자동화의 복잡도다. 캘린더 일정에 따라 일반 알람을 켜고 끄는 핵심 기능은 무료로 배포하고 여러 근무 형태 중 하나를 자동 선택하는 교대 그룹 생성부터 유료로 전환한다.")

make_table(doc, ["상품", "대상", "포함 기능", "과금 방식", "출시 상태"], [
    ("Calarm 무료", "일반 직장인과 기본 모닝콜 사용자", "다중 반복 알람 캘린더 키워드 자동 건너뛰기 휴가 공휴일 알람음 다시 울림 22시 확인", "무료", "출시 기본"),
    ("Calarm 교대 평생 이용권", "교대근무자와 불규칙 출근 근로자", "교대 그룹 생성 근무 형태별 알람 정확한 제목 매칭 충돌 감지 비상 기본 알람", "1회 구매", "첫 유료 상품"),
    ("Calarm 회사 구독", "사업주와 스케줄 담당자", "직원 초대 팀별 캘린더 배정 정책 배포 상태 확인 관리 콘솔", "직원 수 기반 월 구독", "중앙 관리 기능 이후"),
], widths=[3.2, 3.8, 6.6, 2.2, 1.8], font_size=8.1)

heading(doc, "무료 기능 원칙", 2)
bullets(doc, [
    "알람의 기본 신뢰성과 안전 기능에는 과금하지 않는다.",
    "일반 반복 알람과 캘린더 일정에 따른 자동 건너뛰기는 계속 무료로 유지한다.",
    "광고 때문에 알람 화면이나 권한 안내가 방해받지 않도록 초기에는 광고를 넣지 않는 방향을 우선한다.",
    "유료 기능을 구매하지 않아도 이미 만든 일반 알람과 데이터는 정상 작동해야 한다.",
])

heading(doc, "개인 유료 평생 이용권", 2)
body(doc, "첫 유료 상품은 교대 그룹 생성 권한을 한 번 구매하는 비소모성 인앱 상품으로 제안한다. 지속적인 서버 비용이 거의 없는 현재 구조에서는 구독보다 평생 이용권이 사용자 기대와 비용 구조에 맞고 구매 설명도 단순하다.")
make_table(doc, ["항목", "권고안", "검증 방법"], [
    ("확정 정가", "9,900원", "평생 이용권의 기준 가격"),
    ("출시 프로모션", "5,900원", "출시 기간 구매 전환과 초기 사용자 확보"),
    ("출시 전 미리 사용", "상품 소개 후 그룹 생성과 자동 실행 허용", "Play 결제 연결 전 테스트 버전에서 기능 가치 검증"),
    ("복원", "동일 Google Play 계정에서 구매 복원", "재설치와 기기 변경 시 문의율 확인"),
    ("기존 사용자", "유료화 이전 교대 그룹 사용자는 일정 기간 또는 영구 혜택 검토", "부정 리뷰와 신뢰 훼손 방지"),
], widths=[3.0, 6.4, 7.4], font_size=8.5)

callout(doc, "가격 원칙", "정가는 9,900원이고 출시 프로모션 가격은 5,900원이다. 월 구독이 아닌 개인용 평생 이용권이며 실제 결제 연결 전에는 소개 화면을 거쳐 기능을 미리 사용하도록 한다.", fill="FFF4E5", accent="B54708")

heading(doc, "회사 월 구독", 2)
body(doc, "회사 구독은 단순히 공용 캘린더를 읽는 기능만으로 과금하지 않는다. 직원 초대 팀 배정 설정 배포 상태와 같은 중앙 관리 가치가 실제로 구현된 이후 별도 상품으로 출시한다.")
make_table(doc, ["단계", "제안 범위", "가격 가설"], [
    ("회사 파일럿", "공용 캘린더 운영 가이드와 직원별 앱 설치", "무료 또는 개인 이용권으로 검증"),
    ("소규모 팀", "최대 10명 초대 팀 캘린더 배정 기본 상태 확인", "월 19,000원 후보"),
    ("성장 팀", "11명부터 50명 정책 배포와 관리자 역할", "직원당 월 1,500원에서 2,500원 후보"),
    ("맞춤형", "50명 초과 조직과 도입 지원", "별도 견적"),
], widths=[3.0, 8.5, 5.3], font_size=8.7)

heading(doc, "결제 및 비용 원칙", 2)
bullets(doc, [
    "앱 안에서 교대 그룹 같은 디지털 기능을 판매할 때는 원칙적으로 Google Play 결제 시스템을 사용한다.",
    "한국에서 적용되는 수수료와 대체 결제 제도는 출시 시점의 Play Console 정책을 다시 확인한다.",
    "서버 없이 시작하는 개인 평생 이용권은 Play 구매 상태를 기기에서 확인하되 위변조와 환불 처리를 고려한다.",
    "회사 중앙 관리 기능은 인증과 권한과 데이터 동기화가 필요하므로 실제 수요와 매출 가설이 검증된 뒤 서버 비용을 투입한다.",
])

# Marketing
heading(doc, "마케팅 정의", 1)
heading(doc, "포지셔닝", 2)
body(doc, "Calarm은 캘린더 앱도 범용 알람 앱도 아니다. 불규칙한 출근 일정과 기상 시간을 연결하는 근무 스케줄 기반 모닝콜 앱이다. 마케팅에서는 기능의 개수보다 매일 알람을 다시 맞추지 않아도 된다는 결과를 먼저 전달한다.")

heading(doc, "메시지 체계", 2)
make_table(doc, ["대상", "문제 메시지", "해결 메시지", "행동 유도"], [
    ("개인 근로자", "오늘 주간인지 야간인지 확인하고 또 알람을 바꾸시나요", "근무표만 적으면 오늘 알람은 자동으로 선택됩니다", "무료로 첫 알람 연결하기"),
    ("휴가가 잦은 직장인", "휴가인데 평일 알람이 울린 적이 있나요", "휴가 일정이 있는 날만 모닝콜을 자동으로 쉽니다", "캘린더 키워드 연결하기"),
    ("사업주", "근무표를 전달해도 직원이 다시 알람을 맞춰야 하나요", "회사 캘린더와 직원의 기상 시간을 자동으로 연결합니다", "팀 파일럿 시작하기"),
], widths=[3.2, 4.7, 5.5, 3.4], font_size=8.4)

heading(doc, "핵심 카피 후보", 2)
bullets(doc, [
    "근무표는 캘린더에 모닝콜은 Calarm이",
    "주간 야간이 바뀌어도 알람은 자동으로",
    "휴가 날은 쉬고 출근 날은 놓치지 않는 모닝콜",
    "매일 알람을 다시 맞추는 교대근무자에게",
    "회사 스케줄과 내 기상 시간을 한 번에 연결",
])

heading(doc, "Google Play 소개 문안 초안", 2)
callout(doc, "짧은 설명", "근무 일정에 맞춰 필요한 모닝콜을 자동으로 선택하는 캘린더 기반 알람")
body(doc, "상세 설명 첫 문단")
body(doc, "출근 시간이 매일 같지 않다면 알람도 매일 다시 설정해야 했습니다. Calarm은 Google Calendar의 휴가와 근무 형태를 읽어 오늘 필요한 알람만 울립니다. 주간 야간 새벽 근무 시간을 한 번 연결해 두고 앞으로는 캘린더에 근무표만 기록하세요.")

heading(doc, "검색 키워드 후보", 2)
make_table(doc, ["우선순위", "키워드"], [
    ("핵심", "교대근무 알람 교대 알람 근무표 알람 캘린더 알람"),
    ("문제", "야간근무 알람 주간근무 모닝콜 휴가 알람 끄기"),
    ("직군", "간호사 알람 생산직 알람 물류 교대근무 보안 근무 알람"),
    ("브랜드", "Calarm 캘람 캘린더 모닝콜"),
], widths=[3.2, 13.6], font_size=9)

heading(doc, "콘텐츠 소재", 2)
bullets(doc, [
    "주간에서 야간으로 바뀌는 주간에 알람을 매번 수정하는 장면과 자동 선택 장면 비교",
    "휴가 일정을 등록하자 다음 날 알람이 자동으로 쉬는 짧은 시연",
    "사업주가 공용 캘린더를 수정하고 직원별 다른 기상 시간으로 알람이 반영되는 사례",
    "갤럭시 잠금화면에서 알람을 밀어서 해제하고 다시 울림을 사용하는 실제 화면",
    "직군별 근무표 예시를 활용한 블로그와 숏폼 사용법",
])

page_break(doc)

# Go to market and roadmap
heading(doc, "출시 및 성장 전략", 1)
heading(doc, "초기 출시", 2)
numbered(doc, [
    "무료 캘린더 자동 건너뛰기로 설치 장벽을 낮추고 실제 알람 신뢰성을 확보한다.",
    "교대근무 커뮤니티와 직군별 소규모 사용자에게 평생 이용권 가격을 검증한다.",
    "Galaxy S22 Ultra를 포함한 Samsung 기기 중심으로 잠금화면과 절전 설정 가이드를 축적한다.",
    "회사 공용 캘린더 파일럿을 3곳에서 5곳의 소규모 사업장에 적용해 관리자 기능의 필요성을 확인한다.",
    "반복적으로 요청되는 회사 기능만 중앙 관리 제품에 포함한다.",
])

heading(doc, "추천 유입 경로", 2)
make_table(doc, ["경로", "목적", "운영 방식"], [
    ("Google Play 검색", "문제 인지 사용자의 직접 설치", "교대근무 알람과 캘린더 알람 키워드 최적화"),
    ("직군 커뮤니티", "정확한 타깃 검증", "실제 근무표와 알람 자동화 사례 중심"),
    ("숏폼 영상", "기능을 짧게 이해", "일정 입력 전후와 22시 알림과 실제 알람을 20초 내 시연"),
    ("회사 직접 제안", "B2B 문제 검증", "무료 공용 캘린더 파일럿 후 인터뷰"),
    ("추천", "동료 단위 확산", "근무 형태 설정 템플릿 공유 기능은 수요 검증 후 고려"),
], widths=[3.2, 5.2, 8.4], font_size=8.6)

heading(doc, "제품 로드맵", 1)
make_table(doc, ["단계", "목표", "주요 범위", "진입 조건"], [
    ("1단계 안정화", "모닝콜 신뢰 확보", "실기기 알람 잠금화면 캘린더 권한 제조사 진단", "핵심 실패 사례 감소"),
    ("2단계 개인 유료화", "교대 그룹 가치 수익화", "Play Billing 평생 이용권 구매 복원 기능 잠금", "교대 기능 재사용과 구매 의향 확인"),
    ("3단계 회사 파일럿", "공용 캘린더 운영 검증", "회사 도입 가이드 팀별 캘린더 규칙 지원", "3곳 이상 반복 사용"),
    ("4단계 회사 구독", "관리 가치 수익화", "직원 초대 팀 배정 정책 배포 상태 확인", "수동 운영의 명확한 병목 확인"),
    ("5단계 확장", "캘린더 생태계 확대", "Outlook 검토 조직 정책과 보고 기능", "Google 기반 수요와 매출 안정"),
], widths=[2.7, 3.5, 6.3, 4.3], font_size=8.1)

heading(doc, "향후 기능 우선순위", 2)
make_table(doc, ["우선순위", "기능", "이유"], [
    ("높음", "교대 그룹 유료 잠금과 구매 복원", "합의된 첫 유료 상품이며 서버 없이 시작 가능"),
    ("높음", "알람 성공 진단과 기기별 가이드 강화", "상품보다 신뢰가 먼저이며 리뷰에 직접 영향"),
    ("중간", "교대 그룹별 캘린더 지정", "회사와 개인 일정 혼합 오판을 줄임"),
    ("중간", "직원 또는 팀 식별 규칙", "한 공용 캘린더에서 여러 직원 스케줄 지원"),
    ("보류", "회사 중앙 관리 서버", "지속 비용과 개인정보 책임이 생기므로 파일럿 후 결정"),
    ("보류", "Outlook 연동", "Google Calendar 기반 제품 적합성 검증 이후"),
], widths=[2.7, 6.0, 8.1], font_size=8.6)

# Metrics
heading(doc, "성과 측정", 1)
heading(doc, "북극성 지표", 2)
callout(doc, "주간 자동 처리 알람 수", "사용자가 수동으로 다시 설정하지 않고 캘린더 판단에 따라 정상 처리된 알람의 주간 횟수. 제품이 실제 번거로움을 줄였는지 가장 직접적으로 보여준다.", fill=PALE_MINT, accent=MINT)

heading(doc, "핵심 지표", 2)
make_table(doc, ["영역", "지표", "해석"], [
    ("활성화", "설치 후 24시간 이내 첫 알람 저장률", "권한과 초기 설정 장벽"),
    ("핵심 가치", "캘린더 자동 건너뛰기 활성 사용자 비율", "무료 차별 기능 채택"),
    ("유료 가치", "교대 그룹 미리보기 사용자 대비 구매율", "평생 이용권 가격과 설명 적합성"),
    ("유지", "7일과 30일 알람 사용 유지율", "생활 습관에 정착했는지"),
    ("신뢰", "예약 대비 정상 실행률과 캘린더 조회 실패율", "제품의 본질적 품질"),
    ("안전", "권한 미설정과 배터리 제한 상태 비율", "실패 가능성이 있는 기기 규모"),
    ("회사", "파일럿 사업장 주간 활성 직원 비율", "공용 캘린더 운영의 실제 활용"),
    ("사업", "환불률 리뷰 평점 문의 유형", "가격 만족과 장애의 조기 신호"),
], widths=[3.0, 6.3, 7.5], font_size=8.6)

heading(doc, "초기 목표 가설", 2)
bullets(doc, [
    "첫 알람 저장률 60퍼센트 이상",
    "교대 그룹 안내를 본 사용자 중 유료 기능 관심 행동 15퍼센트 이상",
    "평생 이용권 구매 화면 도달 대비 구매율 5퍼센트 이상을 1차 검증 기준으로 사용",
    "정상 권한이 갖춰진 기기에서 예약 대비 알람 실행 성공률 99.5퍼센트 이상 지향",
    "회사 파일럿 참여 직원의 4주 유지율 60퍼센트 이상",
])
body(doc, "위 수치는 확정 목표가 아니라 초기 데이터가 없을 때 사용할 운영 가설이며 실제 설치 모수와 사용자 인터뷰 결과에 따라 조정한다.")

# Risks
heading(doc, "위험요소와 대응", 1)
make_table(doc, ["위험", "영향", "대응 원칙"], [
    ("알람 미실행", "지각과 사용자 신뢰 상실", "정확한 알람 전체화면 배터리 진단과 실기기 회귀 테스트를 최우선"),
    ("캘린더 동기화 지연", "변경된 근무표 미반영", "22시 확인과 실제 알람 시각 재조회 및 동기화 안내"),
    ("근무 제목 충돌", "잘못된 교대 알람 선택", "전체 제목 일치와 충돌 시 그룹 중지 및 경고"),
    ("공용 캘린더 개인정보", "직원 근무와 휴가 정보 노출", "최소 권한 팀별 분리와 공개 캘린더 금지"),
    ("공유 계정 보안", "비밀번호 유출과 퇴사자 접근", "계정 공유 대신 캘린더 공유와 권한 회수"),
    ("유료 전환 반감", "리뷰와 유지율 악화", "무료 핵심 유지와 기존 사용자 정책 사전 고지"),
    ("제조사 백그라운드 제한", "특정 기기에서 지연", "Samsung 우선 검증과 제조사별 진단 문서 축적"),
    ("회사 기능 조기 개발", "서버 비용과 복잡도 증가", "공용 캘린더 파일럿으로 수요 확인 후 투자"),
], widths=[3.4, 5.0, 8.4], font_size=8.3)

heading(doc, "개인정보 원칙", 2)
bullets(doc, [
    "캘린더 읽기 권한의 사용 목적을 알람 판단으로 한정하고 사용자가 선택한 캘린더만 조회한다.",
    "초기 개인 버전은 가능한 한 기기 내부에서 처리하고 캘린더 내용을 별도 서버에 저장하지 않는다.",
    "회사 기능에서 서버가 도입되면 직원 동의 보관 기간 관리자 권한 감사 기록과 탈퇴 및 삭제 절차를 별도로 설계한다.",
    "마케팅 문구에서 절대 실패하지 않는다는 표현을 사용하지 않고 제조사 설정과 동기화 조건을 명확히 알린다.",
])

# Decision summary
heading(doc, "확정 사항과 미결정 사항", 1)
heading(doc, "현재 확정", 2)
make_table(doc, ["항목", "결정"], [
    ("제품 목적", "불규칙 출근 일정에 맞는 모닝콜 자동화"),
    ("무료 범위", "일반 반복 알람과 캘린더 일정에 따른 자동 켜기 끄기"),
    ("개인 유료 시작점", "교대 그룹 생성"),
    ("개인 과금", "초기에는 평생 이용권"),
    ("회사 과금", "중앙 관리 기능이 구현된 시점에 월 구독"),
    ("초기 기술 원칙", "서버와 외부 API 비용 없이 앱과 휴대폰과 연동 캘린더 중심"),
], widths=[4.2, 12.6], font_size=9)

heading(doc, "출시 전 결정 필요", 2)
make_table(doc, ["결정 항목", "추천안", "결정 시점"], [
    ("평생 이용권 정가", "9,900원 확정", "확정"),
    ("출시 할인", "5,900원", "출시 프로모션 기간 적용"),
    ("기존 교대 사용자 처리", "초기 사용자 영구 잠금 해제 또는 충분한 유예", "유료 기능 배포 전"),
    ("구매 전 안내", "자동 켜기 끄기 설명과 일정별 예시 후 구매 유도", "결제 UX 설계 시"),
    ("회사 파일럿 업종", "제조 물류 보안 중 접근 가능한 업종 우선", "개인 유료 출시 후"),
    ("서버 도입", "회사 파일럿에서 중앙 관리 요구가 반복될 때", "파일럿 종료 후"),
], widths=[4.5, 8.2, 4.1], font_size=8.8)

heading(doc, "90일 실행안", 1)
make_table(doc, ["기간", "제품", "마케팅", "사업 검증"], [
    ("1일부터 30일", "알람 실패와 권한 진단 개선 결제 구조 설계", "스토어 문안과 실제 사용 영상 준비", "교대근무자 10명 인터뷰"),
    ("31일부터 60일", "교대 그룹 평생 이용권과 구매 복원 구현", "비공개 테스트와 직군별 콘텐츠", "가격 후보별 구매 의향 확인"),
    ("61일부터 90일", "피드백 수정과 정식 출시 준비", "Google Play 검색 최적화와 초기 리뷰 확보", "소규모 사업장 3곳 공용 캘린더 파일럿 제안"),
], widths=[3.0, 5.2, 4.8, 3.8], font_size=8.4)

heading(doc, "마케팅용 최종 요약", 1)
callout(doc, "Calarm", "출근 시간이 매일 달라도 알람을 매일 다시 맞출 필요는 없습니다. 캘린더에 주간 야간 휴가 같은 일정을 기록하고 기상 시간을 한 번 연결하면 Calarm이 오늘 필요한 모닝콜만 자동으로 울립니다.", fill=PALE_MINT, accent=MINT)

heading(doc, "참고 자료", 1)
p = doc.add_paragraph()
set_run_font(p.add_run("Google Play 결제 정책    "), size=9)
add_hyperlink(p, "Payments Play Console Help", "https://support.google.com/googleplay/android-developer/answer/9858738?hl=en")
p = doc.add_paragraph()
set_run_font(p.add_run("Google Play 서비스 수수료    "), size=9)
add_hyperlink(p, "Service fees Play Console Help", "https://support.google.com/googleplay/android-developer/answer/112622?hl=en")
p = doc.add_paragraph()
set_run_font(p.add_run("Google Calendar 공유 권한    "), size=9)
add_hyperlink(p, "Share your calendar Google Calendar Help", "https://support.google.com/calendar/answer/37082?hl=en")
p = doc.add_paragraph()
set_run_font(p.add_run("Android 캘린더 데이터 공유 설정    "), size=9)
add_hyperlink(p, "Control Google Calendar data sharing", "https://support.google.com/calendar/answer/16089827?hl=en")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
